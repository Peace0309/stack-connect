
from __future__ import annotations

import io
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document
from pptx import Presentation

SUPPORTED = {".pdf", ".docx", ".pptx", ".txt", ".md"}


def parse(file_bytes: bytes, file_name: str) -> str:
    suffix = Path(file_name).suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError(f"Unsupported file type: {suffix}")

    if suffix in {".txt", ".md"}:
        return file_bytes.decode("utf-8", errors="ignore")

    if suffix == ".pdf":
        with fitz.open(stream=file_bytes, filetype="pdf") as doc:
            return "\n".join(page.get_text() for page in doc)

    if suffix == ".docx":
        document = Document(io.BytesIO(file_bytes))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)

    presentation = Presentation(io.BytesIO(file_bytes))
    parts: list[str] = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
    return "\n".join(parts)


def normalise(text: str) -> str:
    return " ".join(text.split())
