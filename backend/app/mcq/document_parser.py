
from __future__ import annotations

import io
from collections import Counter
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document
from pptx import Presentation

SUPPORTED = {".pdf", ".docx", ".pptx", ".txt", ".md"}
def _strip_repeated_lines(pages: list[str]) -> list[str]:
    if len(pages) < 3:
        return pages
    page_lines = [[ln.strip() for ln in p.split("\n")] for p in pages]
    line_counts: Counter[str] = Counter()
    for lines in page_lines:
        line_counts.update(set(ln for ln in lines if ln))
    threshold = max(3, len(pages) // 2)
    boilerplate = {ln for ln, count in line_counts.items() if count >= threshold}
    return ["\n".join(ln for ln in lines if ln not in boilerplate) for lines in page_lines]

def parse(file_bytes: bytes, file_name: str) -> str:
    suffix = Path(file_name).suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError(f"Unsupported file type: {suffix}")

    if suffix in {".txt", ".md"}:
        return file_bytes.decode("utf-8", errors="ignore")

    if suffix == ".pdf":
        with fitz.open(stream=file_bytes, filetype="pdf") as doc:
            pages = [page.get_text() for page in doc]
        pages = _strip_repeated_lines(pages)
        return "\n".join(pages)
    if suffix == ".docx":
        document = Document(io.BytesIO(file_bytes))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)

    presentation = Presentation(io.BytesIO(file_bytes))
    parts: list[str] = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
    return "\n".join(parts)


def normalise(text: str) -> str:
    return " ".join(text.split())
